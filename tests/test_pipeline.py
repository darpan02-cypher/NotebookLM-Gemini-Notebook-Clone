"""Offline tests: fake embedder/LLM, so no model downloads or API keys are needed."""
import hashlib
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from generation import ArtifactService, ChatService  # noqa: E402
from ingestion import SourceIngestor, chunk_text  # noqa: E402
from retrieval.retriever import Retriever  # noqa: E402
from storage import StorageManager  # noqa: E402


class FakeEmbedder:
    """Deterministic bag-of-words hash embedding (dim 64)."""

    def embed(self, texts):
        out = []
        for t in texts:
            v = [0.0] * 64
            for w in t.lower().split():
                v[int(hashlib.md5(w.encode()).hexdigest(), 16) % 64] += 1.0
            norm = sum(x * x for x in v) ** 0.5 or 1.0
            out.append([x / norm for x in v])
        return out


class FakeReranker:
    def score(self, query, passages):
        return [float(sum(w in p.lower() for w in query.lower().split())) for p in passages]


class FakeLLM:
    def generate(self, prompt, system=None):
        return "Answer [1]"


@pytest.fixture
def env(tmp_path):
    storage = StorageManager(tmp_path / "data")
    emb = FakeEmbedder()
    return storage, SourceIngestor(storage, emb), Retriever(storage, emb, FakeReranker())


def test_chunking_size_and_overlap():
    chunks = chunk_text("word " * 1000)
    assert len(chunks) > 3
    assert all(len(c) <= 900 for c in chunks)
    assert chunk_text("") == []


def test_notebook_crud_and_persistence(tmp_path):
    s = StorageManager(tmp_path)
    nb = s.create_notebook("A")
    s.rename_notebook(nb["id"], "B")
    assert StorageManager(tmp_path).get_notebook(nb["id"])["name"] == "B"  # survives "restart"
    s.delete_notebook(nb["id"])
    assert s.list_notebooks() == []


def test_txt_ingest_and_retrieval_both_methods(env, tmp_path):
    storage, ing, retr = env
    nb = storage.create_notebook("T")
    f = tmp_path / "a.txt"
    f.write_text("Photosynthesis converts light into chemical energy. " * 30 + "Mitochondria produce ATP. " * 30)
    res = ing.ingest(nb["id"], "txt", "a.txt", str(f))
    assert res["chunk_count"] > 0
    for method in ("vector", "rerank"):
        out = retr.retrieve(nb["id"], "What do mitochondria produce?", method=method, top_k=3)
        assert out["hits"] and out["hits"][0]["metadata"]["source_name"] == "a.txt"
        assert out["latency_s"] >= 0


def test_pdf_page_metadata(env, tmp_path):
    from reportlab.pdfgen import canvas

    storage, ing, retr = env
    nb = storage.create_notebook("P")
    pdf = tmp_path / "d.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 700, "Page one talks about apples.")
    c.showPage()
    c.drawString(72, 700, "Page two talks about zebras.")
    c.save()
    ing.ingest(nb["id"], "pdf", "d.pdf", str(pdf))
    hit = retr.retrieve(nb["id"], "zebras", method="vector", top_k=1)["hits"][0]
    assert hit["metadata"]["page_or_slide"] == 2


def test_pptx_slide_metadata(env, tmp_path):
    from pptx import Presentation

    storage, ing, retr = env
    nb = storage.create_notebook("S")
    prs = Presentation()
    for title in ["alpha topic", "beta topic"]:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
    path = tmp_path / "d.pptx"
    prs.save(str(path))
    ing.ingest(nb["id"], "pptx", "d.pptx", str(path))
    hit = retr.retrieve(nb["id"], "beta", method="vector", top_k=1)["hits"][0]
    assert hit["metadata"]["page_or_slide"] == 2


def test_notebook_isolation(env, tmp_path):
    storage, ing, retr = env
    a, b = storage.create_notebook("A"), storage.create_notebook("B")
    f = tmp_path / "x.txt"
    f.write_text("secret llama facts")
    ing.ingest(a["id"], "txt", "x.txt", str(f))
    assert retr.retrieve(b["id"], "llama", method="vector")["hits"] == []


def test_failed_ingest_recorded(env):
    storage, ing, _ = env
    nb = storage.create_notebook("F")
    with pytest.raises(ValueError):
        ing.ingest(nb["id"], "url", "bad", "ftp://nope")
    assert storage.get_notebook(nb["id"])["sources"][0]["status"] == "failed"


def test_chat_and_artifacts_persist(env, tmp_path):
    storage, ing, retr = env
    nb = storage.create_notebook("C")
    f = tmp_path / "a.txt"
    f.write_text("Cats are mammals. " * 40)
    ing.ingest(nb["id"], "txt", "a.txt", str(f))
    out = ChatService(storage, retr, FakeLLM()).ask(nb["id"], "What are cats?")
    assert out["citations"][0]["source_name"] == "a.txt"
    assert len(storage.load_chats(nb["id"])) == 2
    art = ArtifactService(storage, FakeLLM())
    rec = art.generate(nb["id"], "quiz")
    assert Path(rec["path"]).exists() and rec["filename"].endswith(".md")
    assert art.list(nb["id"])[0]["type"] == "quiz"
